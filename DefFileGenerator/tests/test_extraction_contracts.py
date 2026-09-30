"""Regression tests for documented extraction and numeric-value contracts."""

import csv
import zipfile
from unittest.mock import MagicMock, patch

import pytest
from defusedxml.common import DTDForbidden
from openpyxl import Workbook

from DefFileGenerator.def_gen import Generator
from DefFileGenerator.extractor import ExtractionError, Extractor


@pytest.mark.parametrize("address", [0, "0"])
def test_address_zero_is_a_valid_register(address):
    rows = [[{"Address": address, "Name": "Status", "Type": "U16"}]]

    mapped = list(Extractor().map_and_clean(rows))

    assert len(mapped) == 1
    assert mapped[0]["Address"] == "0"


@pytest.mark.parametrize("factor", [0, 0.0, "0"])
def test_explicit_zero_factor_takes_precedence_over_gain(factor):
    rows = [[{"Address": 100, "Name": "Constant", "Type": "U16", "Factor": factor, "Gain": 10}]]

    mapped = list(Extractor().map_and_clean(rows))

    assert float(mapped[0]["Factor"]) == 0


def test_numeric_excel_address_zero_survives_extraction(tmp_path):
    source = tmp_path / "zero.xlsx"
    workbook = Workbook()
    workbook.active.append(["Address", "Name", "Type"])
    workbook.active.append([0, "Status", "U16"])
    workbook.active.append([1, "Voltage", "U16"])
    workbook.save(source)

    extractor = Extractor()
    mapped = list(extractor.map_and_clean(extractor.extract_from_excel(str(source))))

    assert [row["Address"] for row in mapped] == ["0", "1"]


@pytest.mark.parametrize("first_address", [0, 10, 100, 1000])
def test_headerless_excel_preserves_first_register(tmp_path, first_address):
    source = tmp_path / "headerless.xlsx"
    workbook = Workbook()
    workbook.active.append([first_address, "Voltage", "U16"])
    workbook.active.append([first_address + 1, "Current", "U16"])
    workbook.save(source)

    extractor = Extractor()
    mapped = list(extractor.map_and_clean(extractor.extract_from_excel(str(source))))

    assert [(row["Address"], row["Name"], row["Type"]) for row in mapped] == [
        (str(first_address), "Voltage", "U16"),
        (str(first_address + 1), "Current", "U16"),
    ]


@pytest.mark.parametrize(
    "declaration",
    ["<!DOCTYPE registers>", '<!DOCTYPE registers SYSTEM "https://example.invalid/vendor.dtd">'],
)
def test_xml_rejects_dtd_declarations_without_entities(tmp_path, declaration):
    source = tmp_path / "dtd.xml"
    source.write_text(
        declaration + "<registers><register><Address>100</Address><Name>Voltage</Name>"
        "<Type>U16</Type></register></registers>",
        encoding="utf-8",
    )

    extractor = Extractor()
    with pytest.raises(DTDForbidden):
        list(extractor.map_and_clean(extractor.extract_from_xml(str(source))))


def test_single_row_pdf_continuation_is_preserved(tmp_path):
    source = tmp_path / "continuation.pdf"
    source.write_bytes(b"%PDF")
    first_page = MagicMock()
    first_page.extract_tables.return_value = [
        [["Address", "Name", "Type"], ["1000", "Voltage", "U16"]]
    ]
    second_page = MagicMock()
    second_page.extract_tables.return_value = [[["1001", "Current", "U16"]]]
    document = MagicMock()
    document.pages = [first_page, second_page]
    extractor = Extractor()

    with patch("DefFileGenerator.extractor.pdfplumber.open") as open_pdf:
        open_pdf.return_value.__enter__.return_value = document
        mapped = list(extractor.map_and_clean(extractor.extract_from_pdf(str(source))))

    assert [(row["Address"], row["Name"]) for row in mapped] == [
        ("1000", "Voltage"),
        ("1001", "Current"),
    ]


def test_sparse_excel_dimensions_are_rejected_before_row_traversal(tmp_path):
    source = tmp_path / "sparse.xlsx"
    workbook = Workbook()
    workbook.active.append(["Address", "Name", "Type"])
    workbook.active.append([100, "Voltage", "U16"])
    workbook.active.cell(1048576, 16384).number_format = "0"
    workbook.save(source)
    assert source.stat().st_size < 10_000

    with (
        patch("openpyxl.worksheet.worksheet.Worksheet.iter_rows") as iterate_rows,
        pytest.raises(ExtractionError, match="maximum expanded worksheet size"),
    ):
        list(Extractor().map_and_clean(Extractor().extract_from_excel(str(source))))

    iterate_rows.assert_not_called()


def test_excel_expanded_cell_limit_is_inclusive(tmp_path, monkeypatch):
    source = tmp_path / "small.xlsx"
    workbook = Workbook()
    workbook.active.append(["Address", "Name", "Type"])
    workbook.active.append([100, "Voltage", "U16"])
    workbook.save(source)
    monkeypatch.setattr("DefFileGenerator.extractor.MAX_EXCEL_SHEET_CELLS", 6)
    extractor = Extractor()

    mapped = list(extractor.map_and_clean(extractor.extract_from_excel(str(source))))

    assert [row["Address"] for row in mapped] == ["100"]


def test_csv_parse_failure_after_valid_row_raises_instead_of_truncating(tmp_path):
    source = tmp_path / "partial.csv"
    source.write_text(
        'Address,Name,Type\n100,First,U16\n101,"' + "A" * 33 + '",U16\n102,Last,U16\n',
        encoding="utf-8",
    )
    original_limit = csv.field_size_limit(32)
    try:
        table = next(Extractor().extract_from_csv(str(source)))
        assert next(table)["Address"] == "100"
        with pytest.raises(ExtractionError, match="CSV parsing failed"):
            next(table)
    finally:
        csv.field_size_limit(original_limit)


def test_gain_reciprocal_is_rounded_after_scale_factor_is_applied():
    rows = [[{"Address": 100, "Name": "Voltage", "Type": "U16", "Gain": 3, "ScaleFactor": 6}]]
    mapped = Extractor().map_and_clean(rows)

    processed = list(Generator().process_rows(mapped))

    assert processed[0]["CoefA"] == "333333.333333"


@pytest.mark.parametrize("length_field,length", [("Length", 10), ("Quantity", 5)])
def test_raw_length_columns_form_a_byte_length_address(length_field, length):
    rows = [[{"Address": 100, "Name": "Payload", "Type": "RAW", length_field: length}]]

    mapped = list(Extractor().map_and_clean(rows))

    assert mapped[0]["Address"] == "100_10"
    processed = list(Generator().process_rows(mapped))
    assert processed[0]["Info2"] == "100_10"


def test_pdf_parse_failure_after_valid_table_raises_instead_of_truncating(tmp_path):
    source = tmp_path / "partial.pdf"
    source.write_bytes(b"%PDF")
    first_page = MagicMock()
    first_page.extract_tables.return_value = [
        [["Address", "Name", "Type"], ["1000", "Voltage", "U16"]]
    ]
    second_page = MagicMock()
    second_page.extract_tables.side_effect = ValueError("invalid page table")
    document = MagicMock()
    document.pages = [first_page, second_page]

    with patch("DefFileGenerator.extractor.pdfplumber.open") as open_pdf:
        open_pdf.return_value.__enter__.return_value = document
        tables = Extractor().extract_from_pdf(str(source))
        assert list(next(tables))[0]["Address"] == "1000"
        with pytest.raises(ExtractionError, match="PDF extraction failed"):
            next(tables)


@pytest.mark.parametrize("limit_kind", ["member", "archive"])
def test_excel_archive_expansion_is_bounded_before_workbook_loading(
    tmp_path, monkeypatch, limit_kind
):
    source = tmp_path / "compressed.xlsx"
    with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("xl/large.xml", b"A" * 33)
        archive.writestr("xl/second.xml", b"B" * 33)
    if limit_kind == "member":
        monkeypatch.setattr("DefFileGenerator.extractor.MAX_EXCEL_ARCHIVE_MEMBER_BYTES", 32)
    else:
        monkeypatch.setattr("DefFileGenerator.extractor.MAX_EXCEL_ARCHIVE_BYTES", 64)

    with (
        patch("DefFileGenerator.extractor.openpyxl.load_workbook") as load_workbook,
        pytest.raises(ExtractionError, match="maximum expanded archive"),
    ):
        list(Extractor().map_and_clean(Extractor().extract_from_excel(str(source))))

    load_workbook.assert_not_called()


def test_excel_archive_expansion_limits_are_inclusive(tmp_path, monkeypatch):
    source = tmp_path / "bounded.xlsx"
    workbook = Workbook()
    workbook.active.append(["Address", "Name", "Type"])
    workbook.active.append([100, "Voltage", "U16"])
    workbook.save(source)
    with zipfile.ZipFile(source) as archive:
        sizes = [member.file_size for member in archive.infolist()]
    monkeypatch.setattr("DefFileGenerator.extractor.MAX_EXCEL_ARCHIVE_BYTES", sum(sizes))
    monkeypatch.setattr("DefFileGenerator.extractor.MAX_EXCEL_ARCHIVE_MEMBER_BYTES", max(sizes))

    extractor = Extractor()
    mapped = list(extractor.map_and_clean(extractor.extract_from_excel(str(source))))

    assert [row["Address"] for row in mapped] == ["100"]


@pytest.mark.parametrize(
    "payload",
    [
        "Address,Name,Type\n100,First,U16\n101,Last,U16".encode("utf-16")[:-1],
        b"\xef\xbb\xbfAddress,Name,Type\n100,First,U16\n101,Last,U1\xff",
    ],
)
def test_bom_declared_csv_rejects_malformed_unicode_without_partial_registers(tmp_path, payload):
    source = tmp_path / "invalid-encoding.csv"
    source.write_bytes(payload)
    extractor = Extractor()

    with pytest.raises(ExtractionError, match="CSV decoding failed"):
        list(extractor.map_and_clean(extractor.extract_from_csv(str(source))))


@pytest.mark.parametrize("encoding", ["utf-16", "utf-8-sig", "cp1252"])
def test_valid_csv_encodings_preserve_register_text(tmp_path, encoding):
    source = tmp_path / "encoded.csv"
    source.write_bytes("Address,Name,Type\n100,Énergie,U16\n".encode(encoding))
    extractor = Extractor()

    mapped = list(extractor.map_and_clean(extractor.extract_from_csv(str(source))))

    assert [(row["Address"], row["Name"], row["Type"]) for row in mapped] == [
        ("100", "Énergie", "U16")
    ]


def test_header_inference_prefers_register_addresses_to_small_numeric_columns():
    table = [
        [1, "Voltage", "U16", 30000, 2],
        [2, "Current", "U16", 30001, 2],
    ]

    inferred = Extractor._infer_table_columns(table)

    assert inferred[3] == "Address"
    assert inferred[2] == "Type"


def test_header_inference_does_not_interpret_numeric_metadata_without_a_type_column():
    table = [[0, "Revision", "Firmware", 100], [1, "Revision", "Firmware", 101]]

    assert Extractor._infer_table_columns(table) is None


def test_header_inference_does_not_guess_between_multiple_small_numeric_columns():
    table = [[1, "Voltage", "U16", 10], [2, "Current", "U16", 11]]

    assert Extractor._infer_table_columns(table) is None


def test_excel_custom_header_mapping_remains_authoritative(tmp_path):
    source = tmp_path / "custom.xlsx"
    workbook = Workbook()
    workbook.active.append(["slot", "label", "encoding"])
    workbook.active.append([0, "Voltage", "U16"])
    workbook.save(source)
    extractor = Extractor(mapping={"Address": "slot", "Name": "label", "Type": "encoding"})

    mapped = list(extractor.map_and_clean(extractor.extract_from_excel(str(source))))

    assert [(row["Address"], row["Name"]) for row in mapped] == [("0", "Voltage")]
