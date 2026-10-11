"""
Tests for Multi-Sheet Excel (XLSX) support, Modbus Function Codes, and Action Codes.
Validates:
1. Multi-sheet Excel workbook extraction (all sheets merged vs specific sheet).
2. Resilient header detection against data cells containing words like 'Forward' (containing 'rw').
3. Universal Function Code mapping (FC01-FC04, FC06, FC16, HR, IR, DI, DO, French/English headers).
4. Official WebdynSunPM Action code normalization (0=Reserve, 1=Write/Command, 4=Telemetry, 9=Alarm/Status).
"""

import openpyxl
import pytest

from DefFileGenerator.def_gen import (
    MODBUS_COIL,
    MODBUS_DISCRETE,
    MODBUS_HOLDING,
    MODBUS_INPUT,
    Generator,
)
from DefFileGenerator.extractor import Extractor
from DefFileGenerator.knowledge_miner import parse_function_code_from_text


@pytest.fixture
def multisheet_workbook_path(tmp_path):
    """Creates a temporary multi-sheet Excel workbook with diverse Modbus tables."""
    wb = openpyxl.Workbook()
    # Sheet 1: Inverters (Holding registers)
    ws1 = wb.active
    ws1.title = "Inverters"
    ws1.append(["Addr", "Name", "Type", "Unit", "Access", "Function Code"])
    ws1.append([100, "Active Power", "float32", "kW", "RO", "FC03"])
    ws1.append([102, "Reactive Power", "float32", "kvar", "RO", "FC03"])
    ws1.append([104, "Set Power Limit", "uint16", "%", "RW", "FC16"])

    # Sheet 2: Meters (Input registers, with forward energy)
    ws2 = wb.create_sheet(title="Meters")
    ws2.append(["Register Settings (Function Code: 0x04 - Read Input Registers)"])
    ws2.append(["Addr", "Equipment", "Name", "Unit", "Type", "Notes"])
    # Row with 'Total Forward Active Energy' containing 'forward' (substring 'rw')
    ws2.append([10000, "1#Meter", "Total Forward Active Energy", "kWh", "float", "Measurement"])
    ws2.append([10002, "1#Meter", "Total Forward Reactive Energy", "kvarh", "float", None])
    ws2.append([10004, "1#Meter", "Grid Frequency", "Hz", "float", None])

    # Sheet 3: Alarms & Status (Discrete inputs & Coils)
    ws3 = wb.create_sheet(title="Alarms")
    ws3.append(["Address", "Signal Name", "Data Format", "FC", "R/W"])
    ws3.append([200, "Inverter Grid Overvoltage Alarm", "U16", "DI", "RO"])
    ws3.append([201, "Emergency Trip Breaker", "U16", "DO", "RW"])
    ws3.append([202, "Reserved Point", "U16", "HR", "RO"])

    file_path = str(tmp_path / "multisheet_modbus_device.xlsx")
    wb.save(file_path)
    wb.close()
    return file_path


def test_multisheet_excel_all_sheets_extraction(multisheet_workbook_path):
    """Verifies that all sheets are extracted and merged when sheet_name is None."""
    ext = Extractor()
    raw_tables = list(ext.extract_from_excel(multisheet_workbook_path, sheet_name=None))
    assert len(raw_tables) == 3, f"Expected 3 sheets, got {len(raw_tables)}"

    cleaned = list(ext.map_and_clean(raw_tables))
    # 3 in Inverters + 3 in Meters + 3 in Alarms = 9 registers
    assert len(cleaned) == 9, f"Expected 9 cleaned registers, got {len(cleaned)}"

    names = [r["Name"] for r in cleaned]
    assert "Active Power" in names
    assert "Total Forward Active Energy" in names
    assert "Inverter Grid Overvoltage Alarm" in names


def test_multisheet_excel_single_sheet_selection(multisheet_workbook_path):
    """Verifies that only the targeted sheet is extracted when sheet_name is specified."""
    ext = Extractor()
    raw_tables = list(ext.extract_from_excel(multisheet_workbook_path, sheet_name="Meters"))
    assert len(raw_tables) == 1

    cleaned = list(ext.map_and_clean(raw_tables))
    assert len(cleaned) == 3
    addresses = [r["Address"] for r in cleaned]
    assert addresses == ["10000", "10002", "10004"]
    assert cleaned[0]["Name"] == "Total Forward Active Energy"
    assert cleaned[0]["Type"] == "F32"
    assert cleaned[0]["Unit"] == "kWh"


def test_forward_energy_not_misidentified_as_secondary_header():
    """Confirms that data rows containing 'Forward' or numbers are never treated as headers."""
    ext = Extractor()
    data_row_1 = [10000, "1#Meter", "Total Forward Active Energy", "kWh", "float"]
    data_row_2 = ["10000", "1#Meter", "Total Forward Active Energy", "kWh", "float"]
    header_row = ["Addr", "Equipment", "Name", "Unit", "Type"]

    assert not ext._is_secondary_header_row(data_row_1)
    assert not ext._is_secondary_header_row(data_row_2)
    assert ext._is_secondary_header_row(header_row)


def test_modbus_function_codes_resolution():
    """Validates that various function code representations map to Webdyn Info1 codes."""
    gen = Generator()
    assert gen.register_type_map["fc01"] == MODBUS_COIL
    assert gen.register_type_map["fc02"] == MODBUS_DISCRETE
    assert gen.register_type_map["fc03"] == MODBUS_HOLDING
    assert gen.register_type_map["fc04"] == MODBUS_INPUT
    assert gen.register_type_map["fc06"] == MODBUS_HOLDING
    assert gen.register_type_map["fc16"] == MODBUS_HOLDING
    assert gen.register_type_map["6"] == MODBUS_HOLDING
    assert gen.register_type_map["16"] == MODBUS_HOLDING
    assert gen.register_type_map["hr"] == MODBUS_HOLDING
    assert gen.register_type_map["ir"] == MODBUS_INPUT
    assert gen.register_type_map["di"] == MODBUS_DISCRETE
    assert gen.register_type_map["do"] == MODBUS_COIL

    # Text banner parsing
    assert parse_function_code_from_text("Function Code: 0x04 - Read Input Registers") == "4"
    assert parse_function_code_from_text("FC 03 Holding Registers") == "3"
    assert parse_function_code_from_text("Code Fonction 02 - Discrete") == "2"
    assert parse_function_code_from_text("FC 01 Read Coils") == "1"
    assert parse_function_code_from_text("Function Code: 16 (Write Multiple Registers)") == "3"


def test_modbus_action_codes_end_to_end(multisheet_workbook_path):
    """Verifies that Action codes are correctly assigned and processed by Webdyn Generator."""
    ext = Extractor()
    gen = Generator()

    raw_tables = list(ext.extract_from_excel(multisheet_workbook_path))
    cleaned = list(ext.map_and_clean(raw_tables))
    processed = list(gen.process_rows(cleaned))

    assert len(processed) == 9

    row_map = {r["Name"]: r for r in processed}

    # 1. Active Power: Telemetry -> Action 4
    assert row_map["Active Power"]["Action"] == "4"
    assert row_map["Active Power"]["Info1"] == "3"

    # 2. Set Power Limit: Write command with RW / FC16 -> Action 1
    assert row_map["Set Power Limit"]["Action"] == "1"
    assert row_map["Set Power Limit"]["Info1"] == "3"

    # 3. Total Forward Active Energy: Telemetry -> Action 4
    assert row_map["Total Forward Active Energy"]["Action"] == "4"
    assert row_map["Total Forward Active Energy"]["Info1"] == "4"

    # 4. Inverter Grid Overvoltage Alarm: Alarm -> Action 9
    assert row_map["Inverter Grid Overvoltage Alarm"]["Action"] == "9"
    assert row_map["Inverter Grid Overvoltage Alarm"]["Info1"] == "2"

    # 5. Emergency Trip Breaker: Command -> Action 1
    assert row_map["Emergency Trip Breaker"]["Action"] == "1"
    assert row_map["Emergency Trip Breaker"]["Info1"] == "1"

    # 6. Reserved Point: Reserve -> Action 0
    assert row_map["Reserved Point"]["Action"] == "0"
